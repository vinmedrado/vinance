from __future__ import annotations

import asyncio
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import TYPE_CHECKING, Any, Iterable

from openpyxl import load_workbook

if TYPE_CHECKING:
    from sqlalchemy.ext.asyncio import AsyncSession


PROJECT_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_XLSX_PATH = PROJECT_ROOT / "data" / "imports" / "b3.xlsx"

SHEET_MARKET_ALIASES = {
    "fiis": "fii",
    "fii": "fii",
    "acoes": "acoes",
    "ações": "acoes",
    "acao": "acoes",
    "ação": "acoes",
    "etf": "etf",
    "etfs": "etf",
    "bdr": "bdr",
    "bdrs": "bdr",
}

TICKER_HEADERS = {
    "ticker",
    "codigo",
    "código",
    "cod",
    "ativo",
    "papel",
    "mercados",
    "symbol",
}

NAME_HEADERS = {
    "nome",
    "name",
    "empresa",
    "companhia",
    "razao social",
    "razão social",
    "descricao",
    "descrição",
}

ACTIVE_HEADERS = {"active", "is_active", "ativo", "ativa", "status"}
FALSE_VALUES = {"false", "falso", "nao", "não", "n", "0", "inativo", "inativa", "inactive"}
TRUE_VALUES = {"true", "verdadeiro", "sim", "s", "1", "ativo", "ativa", "active"}


@dataclass(frozen=True)
class CatalogRow:
    ticker: str
    market: str
    name: str
    is_active: bool = True


def _strip_accents(value: str) -> str:
    return "".join(
        char for char in unicodedata.normalize("NFKD", value) if not unicodedata.combining(char)
    )


def _norm_text(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", " ", text)
    return text.lower()


def _norm_key(value: Any) -> str:
    return _strip_accents(_norm_text(value))


def normalize_sheet_market(sheet_name: str) -> str | None:
    raw = _norm_text(sheet_name)
    no_accents = _strip_accents(raw)
    market = SHEET_MARKET_ALIASES.get(raw) or SHEET_MARKET_ALIASES.get(no_accents)
    if not market:
        return None
    return str(market).strip().lower()


def _find_column(headers: list[Any], candidates: set[str]) -> int | None:
    normalized_candidates = {_norm_key(item) for item in candidates}
    for index, header in enumerate(headers):
        if _norm_key(header) in normalized_candidates:
            return index
    return None


def _parse_active(value: Any) -> bool:
    if value is None or str(value).strip() == "":
        return True
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    normalized = _norm_key(value)
    if normalized in FALSE_VALUES:
        return False
    if normalized in TRUE_VALUES:
        return True
    return True


def _clean_ticker(value: Any) -> str | None:
    if value is None:
        return None
    raw = str(value).strip().upper()
    if not raw:
        return None
    # Keep B3 ticker symbols only; discard accidental totals/labels from spreadsheets.
    ticker = re.sub(r"[^A-Z0-9]", "", raw)
    if not ticker or ticker in {"CODIGO", "TICKER", "MERCADOS", "ATIVO"}:
        return None
    return ticker


def _clean_name(value: Any, fallback_ticker: str) -> str:
    name = str(value or "").strip()
    return name[:255] if name else fallback_ticker


def extract_catalog_rows(xlsx_path: Path = DEFAULT_XLSX_PATH) -> list[CatalogRow]:
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {xlsx_path}")

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    rows: list[CatalogRow] = []
    seen: set[tuple[str, str]] = set()

    for sheet_name in workbook.sheetnames:
        market = normalize_sheet_market(sheet_name)
        if market is None:
            continue

        worksheet = workbook[sheet_name]
        iterator = worksheet.iter_rows(values_only=True)
        headers = list(next(iterator, []) or [])
        ticker_col = _find_column(headers, TICKER_HEADERS)
        name_col = _find_column(headers, NAME_HEADERS)
        active_col = _find_column(headers, ACTIVE_HEADERS)

        if ticker_col is None:
            # B3 imports commonly place ticker in the first column even with non-standard headers.
            ticker_col = 0

        for row in iterator:
            ticker = _clean_ticker(row[ticker_col] if ticker_col < len(row) else None)
            if not ticker:
                continue
            key = (ticker, market)
            if key in seen:
                continue
            seen.add(key)

            raw_name = row[name_col] if name_col is not None and name_col < len(row) else None
            raw_active = row[active_col] if active_col is not None and active_col < len(row) else None
            rows.append(
                CatalogRow(
                    ticker=ticker,
                    market=market,
                    name=_clean_name(raw_name, ticker),
                    is_active=_parse_active(raw_active),
                )
            )

    return rows


async def import_catalog_rows(session: "AsyncSession", rows: Iterable[CatalogRow]) -> dict[str, Any]:
    from sqlalchemy import select

    from backend.app.catalog.models import AssetCatalog
    from backend.app.catalog.schemas import default_currency_for_market, normalize_market, normalize_ticker

    inserted = 0
    updated = 0
    skipped = 0
    by_market: dict[str, int] = {}
    now = datetime.now(timezone.utc)

    for row in rows:
        if not row.ticker or not row.market:
            skipped += 1
            continue

        market = normalize_market(row.market)
        ticker = normalize_ticker(row.ticker)
        by_market[market] = by_market.get(market, 0) + 1

        result = await session.execute(
            select(AssetCatalog).where(AssetCatalog.ticker == ticker, AssetCatalog.market == market)
        )
        existing = result.scalar_one_or_none()

        if existing is None:
            session.add(
                AssetCatalog(
                    ticker=ticker,
                    name=row.name or ticker,
                    market=market,
                    currency=default_currency_for_market(market),
                    source="data/imports/b3.xlsx",
                    is_active=row.is_active,
                    updated_at=now,
                )
            )
            inserted += 1
            continue

        existing.name = row.name or existing.name or ticker
        existing.currency = existing.currency or default_currency_for_market(market)
        existing.source = existing.source or "data/imports/b3.xlsx"
        existing.is_active = row.is_active
        if hasattr(existing, "updated_at"):
            existing.updated_at = now
        updated += 1

    await session.commit()
    return {
        "status": "SUCCESS",
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
        "total": inserted + updated,
        "by_market": dict(sorted(by_market.items())),
    }


async def import_b3_catalog(xlsx_path: Path = DEFAULT_XLSX_PATH) -> dict[str, Any]:
    from backend.app.core.database import AsyncSessionLocal

    rows = extract_catalog_rows(xlsx_path)
    async with AsyncSessionLocal() as session:
        return await import_catalog_rows(session, rows)


def _print_summary(result: dict[str, Any]) -> None:
    by_market = result.get("by_market", {}) or {}
    print("Resumo importação B3 -> asset_catalog")
    print(f"FIIs: {by_market.get('fii', 0)}")
    print(f"AÇÕES: {by_market.get('acoes', 0)}")
    print(f"ETFs: {by_market.get('etf', 0)}")
    print(f"BDRs: {by_market.get('bdr', 0)}")
    print(f"Inseridos: {result.get('inserted', 0)}")
    print(f"Atualizados: {result.get('updated', 0)}")
    print(f"Ignorados: {result.get('skipped', 0)}")


def main() -> None:
    result = asyncio.run(import_b3_catalog())
    _print_summary(result)


if __name__ == "__main__":
    try:
        main()
    finally:
        try:
            from backend.app.core.database import close_database

            asyncio.run(close_database())
        except RuntimeError:
            pass
