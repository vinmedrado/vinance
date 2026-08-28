from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from backend.app.market.scripts.import_b3_catalog import (
    CatalogRow,
    extract_catalog_rows,
    normalize_sheet_market,
)


def _save_workbook(path: Path) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    fiis = workbook.create_sheet("fiis")
    fiis.append(["Mercados", "Nome", "is_active"])
    fiis.append(["MXRF11", "Maxi Renda", True])
    fiis.append(["MXRF11", "Duplicado", True])

    acoes = workbook.create_sheet("ações")
    acoes.append(["Código", "Nome", "active"])
    acoes.append(["PETR4", "Petrobras PN", "sim"])
    acoes.append(["VALE3", "Vale ON", "não"])

    etfs = workbook.create_sheet("etfs")
    etfs.append(["ticker", "name"])
    etfs.append(["BOVA11", "iShares BOVA"])

    bdrs = workbook.create_sheet("bdr")
    bdrs.append(["Código", "Nome"])
    bdrs.append(["AAPL34", "Apple BDR"])

    ignored = workbook.create_sheet("demo")
    ignored.append(["Código"])
    ignored.append(["DEMO11"])

    workbook.save(path)


def test_normalize_sheet_market_accepts_required_aliases() -> None:
    assert normalize_sheet_market("fiis") == "fii"
    assert normalize_sheet_market("ações") == "acoes"
    assert normalize_sheet_market("acoes") == "acoes"
    assert normalize_sheet_market("etf") == "etf"
    assert normalize_sheet_market("etfs") == "etf"
    assert normalize_sheet_market("bdr") == "bdr"
    assert normalize_sheet_market("bdrs") == "bdr"


def test_extract_catalog_rows_detects_required_sheets_and_deduplicates(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "b3.xlsx"
    _save_workbook(xlsx_path)

    rows = extract_catalog_rows(xlsx_path)

    assert rows == [
        CatalogRow(ticker="MXRF11", market="fii", name="Maxi Renda", is_active=True),
        CatalogRow(ticker="PETR4", market="acoes", name="Petrobras PN", is_active=True),
        CatalogRow(ticker="VALE3", market="acoes", name="Vale ON", is_active=False),
        CatalogRow(ticker="BOVA11", market="etf", name="iShares BOVA", is_active=True),
        CatalogRow(ticker="AAPL34", market="bdr", name="Apple BDR", is_active=True),
    ]


def test_real_b3_import_file_extracts_assets() -> None:
    xlsx_path = Path("data/imports/b3.xlsx")
    rows = extract_catalog_rows(xlsx_path)
    markets = {row.market for row in rows}

    assert {"fii", "acoes", "etf", "bdr"}.issubset(markets)
    assert len(rows) > 10
    assert all(row.ticker for row in rows)
